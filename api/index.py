import os
import io
import re
import json
import base64
import hashlib
import mimetypes
import threading
from datetime import datetime, timedelta

import requests
from cachetools import TTLCache
from flask import Flask, request, jsonify, render_template, Response, stream_with_context
from groq import Groq, RateLimitError, APIStatusError, APIConnectionError

#DEFINISIKAN BASE_DIR
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
#inisialisasi flask app
app = Flask(
    __name__,
    template_folder=os.path.join(BASE_DIR, 'templates'),
    static_folder=os.path.join(BASE_DIR, 'static'),
    static_url_path='/static'
)

try:
    from pypdf import PdfReader
    PYPDF_AVAILABLE = True
except ImportError:
    try:
        from PyPDF2 import PdfReader
        PYPDF_AVAILABLE = True
    except ImportError:
        PYPDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

#KONFIGURASI
def _env_int(name, default):
    try:
        return int(os.environ.get(name, default))
    except (TypeError, ValueError):
        return default

GROQ_API_KEY = os.environ.get("GROQ_API_KEY")
TEXT_MODEL = os.environ.get("GROQ_MODEL") or "openai/gpt-oss-120b"
VISION_MODEL = os.environ.get("GROQ_VISION_MODEL") or "qwen/qwen3.8-27b"

#max_retries=1:default SDK (2x retry + backoff) membuat request 429 menggantung lama di serverless
client = Groq(api_key=GROQ_API_KEY, max_retries=1) if GROQ_API_KEY else None

NASA_API_KEY = os.environ.get("NASA_API_KEY")

#sliding window & pembatasan ukuran payload
MAX_HISTORY_MESSAGES = _env_int("MAX_HISTORY_MESSAGES", 6)          #pesan terakhir yang dikirim ke Groq
MAX_HISTORY_CHARS_PER_MSG = _env_int("MAX_HISTORY_CHARS_PER_MSG", 3000)
MAX_FILE_CHARS = _env_int("MAX_FILE_CHARS", 12000)                  #batas teks per file
MAX_TOTAL_FILE_CHARS = _env_int("MAX_TOTAL_FILE_CHARS", 24000)      #batas total teks semua file per request
MAX_IMAGES_PER_REQUEST = _env_int("MAX_IMAGES_PER_REQUEST", 3)      #batas 3 gambar
MAX_OUTPUT_TOKENS = _env_int("MAX_OUTPUT_TOKENS", 2048)

#cache in-memory
CACHE_TTL_SECONDS = _env_int("CACHE_TTL_SECONDS", 3600)
CACHE_MAX_ENTRIES = _env_int("CACHE_MAX_ENTRIES", 300)
CACHE_MAX_ANSWER_CHARS = _env_int("CACHE_MAX_ANSWER_CHARS", 20000)


NASA_API_KEY = os.environ.get("NASA_API_KEY")

NASA_EPIC_METADATA_URL = "https://api.nasa.gov/EPIC/api/natural/date/{date}"
NASA_EPIC_IMAGE_URL = "https://api.nasa.gov/EPIC/archive/natural/{yyyy}/{mm}/{dd}/png/{name}.png"

EPIC_LOOKBACK_DAYS = 10
EPIC_REQUEST_TIMEOUT = 10

EARTH_IMAGE_KEYWORDS = [
    "bumi hari ini", "gambar bumi hari ini", "foto bumi hari ini",
    "tampilkan bumi", "lihat bumi", "bumi dari luar angkasa",
    "foto bumi terbaru", "gambar bumi terbaru", "gambar bumi",
    "foto bumi", "epic earth", "nasa earth", "earth today",
    "today's earth", "today earth", "earth from space",
    "view of earth", "planet bumi hari ini",
]

class EpicFetchError(Exception):
    pass

def is_earth_image_intent(message: str) -> bool:
    if not message:
        return False
    text = message.lower().strip()
    return any(keyword in text for keyword in EARTH_IMAGE_KEYWORDS)

def _fetch_epic_metadata_for_date(date_obj):
    date_str = date_obj.strftime("%Y-%m-%d")
    url = NASA_EPIC_METADATA_URL.format(date=date_str)

    try:
        resp = requests.get(url, params={"api_key": NASA_API_KEY}, timeout=EPIC_REQUEST_TIMEOUT)
    except requests.exceptions.Timeout:
        raise EpicFetchError("Waktu koneksi ke server NASA habis (timeout). Silakan coba lagi.")
    except requests.exceptions.ConnectionError:
        raise EpicFetchError("Tidak bisa terhubung ke server NASA. Periksa koneksi internet Anda.")
    except requests.exceptions.RequestException as e:
        raise EpicFetchError(f"Gagal menghubungi NASA EPIC API: {e}")

    if resp.status_code == 429:
        raise EpicFetchError("Batas permintaan (rate limit) API NASA sudah tercapai. Coba lagi nanti.")
    if resp.status_code == 403:
        raise EpicFetchError("API Key NASA tidak valid, belum diset, atau ditolak server.")
    if resp.status_code != 200:
        raise EpicFetchError(f"NASA EPIC API mengembalikan status tidak terduga: {resp.status_code}.")

    try:
        data = resp.json()
    except ValueError:
        raise EpicFetchError("Respons dari NASA EPIC API tidak valid (bukan JSON).")

    if not isinstance(data, list):
        return []
    return data

def _build_epic_result(item, date_obj):
    image_name = item.get("image", "epic_earth")
    caption = item.get("caption") or "Tidak ada caption dari NASA untuk gambar ini."
    date_taken = item.get("date") or date_obj.strftime("%Y-%m-%d %H:%M:%S")

    yyyy, mm, dd = date_obj.strftime("%Y"), date_obj.strftime("%m"), date_obj.strftime("%d")
    image_url = (
        NASA_EPIC_IMAGE_URL.format(yyyy=yyyy, mm=mm, dd=dd, name=image_name)
        + f"?api_key={NASA_API_KEY}"
    )

    return {
        "ok": True,
        "image_name": image_name,
        "image_url": image_url,
        "caption": caption,
        "date": date_taken,
        "date_only": date_obj.strftime("%Y-%m-%d"),
    }

def find_latest_epic_image():
    today = datetime.utcnow().date()
    for offset in range(EPIC_LOOKBACK_DAYS + 1):
        check_date = today - timedelta(days=offset)
        data = _fetch_epic_metadata_for_date(check_date)
        if data:
            return _build_epic_result(data[0], check_date)

    raise EpicFetchError(
        f"Tidak ditemukan gambar Bumi NASA EPIC dalam {EPIC_LOOKBACK_DAYS} hari terakhir."
    )

def get_earth_today_payload():
    try:
        return find_latest_epic_image()
    except EpicFetchError as e:
        return {"ok": False, "error": str(e)}
    except Exception as e:
        return {"ok": False, "error": f"Terjadi kesalahan tak terduga saat mengambil data NASA: {e}"}
'''
PENTING (Groq prompt caching): prompt ini HARUS statis. Jangan sisipkan timestamp, tanggal,
nama user, ID acak, atau nilai dinamis apa pun. Satu karakter berbeda = cache miss
'''
SYSTEM_PROMPT = (
    "Kamu adalah ZUPIUM, asisten AI cerdas, serba bisa, ramah, dan to the point. "
    "Jawab dalam Bahasa Indonesia kecuali user memakai bahasa lain. "
    "Gunakan markdown (heading, list, tabel, **bold**) dan code block ```bahasa``` "
    "saat menjelaskan kode atau tugas supaya rapi dibaca. "
    "Jika user mengirim gambar, deskripsikan dan analisis isinya secara mendetail. "
    "Jika user mengirim file/dokumen (seperti PDF, Word, Excel, PowerPoint, Jupyter Notebook, "
    "atau berbagai format kode pemrograman), baca isinya dengan seksama dan bantu jawab pertanyaan, "
    "perbaiki error, kerjakan tugas, atau buat rangkuman sesuai instruksi user. "
    "Untuk rumus/persamaan matematika, selalu gunakan notasi LaTeX dengan delimiter "
    "\\(... \\) untuk inline dan \\[... \\] untuk blok (atau $$...$$ untuk blok), "
    "termasuk saat memakai lingkungan seperti \\begin{aligned}...\\end{aligned}. "
    "Jangan pernah menampilkan simbol LaTeX mentah di luar delimiter tersebut. "
    "pencipta kamu adalah orang yang biasa di panggil 'GANDHI'"
)

def extract_text_from_pdf(raw_bytes: bytes) -> str:
    if not PYPDF_AVAILABLE:
        return "(Library 'pypdf' belum terpasang di server. Jalankan 'pip install pypdf' di terminal server.)"
    try:
        reader = PdfReader(io.BytesIO(raw_bytes))
        extracted = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted.append(f"--- Halaman {idx + 1} ---\n{page_text.strip()}")
        if not extracted:
            return "(PDF ini tidak memiliki teks yang bisa diekstrak langsung, kemungkinan hasil scan/gambar.)"
        return "\n\n".join(extracted)
    except Exception as e:
        return f"(Gagal membaca isi dokumen PDF: {e})"

def extract_text_from_docx(raw_bytes: bytes) -> str:
    if not DOCX_AVAILABLE:
        return "(Library 'python-docx' belum terpasang di server. Jalankan 'pip install python-docx' di terminal server.)"
    try:
        doc = Document(io.BytesIO(raw_bytes))
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for t_idx, table in enumerate(doc.tables):
            table_lines = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append(" | ".join(cells))
            if table_lines:
                parts.append(f"[Tabel {t_idx + 1}]\n" + "\n".join(table_lines))
        return "\n\n".join(parts) or "(Dokumen Word kosong)"
    except Exception as e:
        return f"(Gagal membaca file Word: {e})"

def extract_text_from_excel(raw_bytes: bytes) -> str:
    if not OPENPYXL_AVAILABLE:
        return "(Library 'openpyxl' belum terpasang di server. Jalankan 'pip install openpyxl' di terminal server.)"
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        sheets_data = []
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows_data = []
            for row in ws.iter_rows(values_only=True, max_row=150):
                if any(cell is not None for cell in row):
                    cells = [str(cell) if cell is not None else "" for cell in row]
                    while cells and cells[-1] == "":
                        cells.pop()
                    if cells:
                        rows_data.append(" | ".join(cells))
            if rows_data:
                sheets_data.append(f"[Sheet: {sheetname}]\n" + "\n".join(rows_data))
        return "\n\n".join(sheets_data) or "(Spreadsheet Excel kosong)"
    except Exception as e:
        return f"(Gagal membaca file Excel: {e})"

def extract_text_from_pptx(raw_bytes: bytes) -> str:
    if not PPTX_AVAILABLE:
        return "(Library 'python-pptx' belum terpasang di server. Jalankan 'pip install python-pptx' di terminal server.)"
    try:
        prs = Presentation(io.BytesIO(raw_bytes))
        slides_text = []
        for idx, slide in enumerate(prs.slides):
            lines = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
            if lines:
                slides_text.append(f"[Slide {idx + 1}]\n" + "\n".join(lines))
        return "\n\n".join(slides_text) or "(File PowerPoint kosong)"
    except Exception as e:
        return f"(Gagal membaca file PowerPoint: {e})"

def extract_text_from_ipynb(raw_bytes: bytes) -> str:
    try:
        content = raw_bytes.decode("utf-8", errors="replace")
        nb = json.loads(content)
        cells = nb.get("cells", [])
        output = []
        for idx, cell in enumerate(cells):
            cell_type = cell.get("cell_type", "code")
            src = "".join(cell.get("source", []))
            if not src.strip():
                continue
            if cell_type == "markdown":
                output.append(f"[Sel {idx + 1} - Markdown]\n{src.strip()}")
            elif cell_type == "code":
                output.append(f"[Sel {idx + 1} - Code]\n```python\n{src.strip()}\n```")
        return "\n\n".join(output) or "(Notebook kosong)"
    except Exception as e:
        return f"(Gagal membaca Jupyter Notebook: {e})"

def try_decode_as_text(raw_bytes: bytes):
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = raw_bytes.decode(enc)
            if "\x00" not in text[:1024]:
                return text
        except UnicodeDecodeError:
            continue
    return None

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".svg"}

CODE_AND_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml",
    ".env", ".ini", ".toml", ".log", ".conf", ".cfg", ".sql", ".tex", ".rtf",
    ".py", ".pyw", ".js", ".mjs", ".cjs", ".jsx", ".ts", ".tsx",
    ".html", ".htm", ".css", ".scss", ".sass", ".less", ".php",
    ".c", ".cpp", ".cc", ".cxx", ".h", ".hpp", ".cs",
    ".java", ".kt", ".kts", ".swift", ".go", ".rs", ".rb", ".dart",
    ".sh", ".bash", ".zsh", ".bat", ".cmd", ".ps1",
    ".r", ".m", ".lua", ".asm", ".v", ".sv", ".pl", ".pm", ".groovy", ".gradle"
}

MAX_FILE_SIZE = 15 * 1024 * 1024

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/health")
def health():
    return jsonify({"status": "ok", "model": TEXT_MODEL, "vision_model": VISION_MODEL, "configured": bool(GROQ_API_KEY), "cache_entries": len(_response_cache)})

@app.route("/api/earth-today")
def earth_today():
    payload = get_earth_today_payload()
    status_code = 200 if payload.get("ok") else 502
    return jsonify(payload), status_code

def build_user_content(user_message, attachments):
    has_image = any(a["kind"] == "image" for a in attachments)

    if not has_image:
        extra_text = ""
        for a in attachments:
            if a["kind"] == "text":
                extra_text += f"\n\n--- Isi file: {a['name']} ---\n{a['text']}\n--- akhir file ---\n"
        return user_message + extra_text, False

    content = []
    if user_message:
        content.append({"type": "text", "text": user_message})

    extra_text = ""
    for a in attachments:
        if a["kind"] == "image":
            content.append({
                "type": "image_url",
                "image_url": {"url": f"data:{a['mime']};base64,{a['data_b64']}"}
            })
        elif a["kind"] == "text":
            extra_text += f"\n\n--- Isi file: {a['name']} ---\n{a['text']}\n--- akhir file ---\n"

    if extra_text:
        if content and content[0]["type"] == "text":
            content[0]["text"] += extra_text
        else:
            content.insert(0, {"type": "text", "text": extra_text.strip()})

    if not content:
        content.append({"type": "text", "text": "(tidak ada teks)"})

    return content, True


#SOLUSI 2:SLIDING WINDOW UNTUK RIWAYAT CHA
#pesan assistant berikut adalah pesan error/placeholder dari frontend, tidak berguna sebagai konteks
_NOISE_PREFIXES = ("\u26a0\ufe0f", "_Tidak ada respons", "_Permintaan dihentikan")

def _content_to_text(content):
    """Ambil teksnya saja. Bagian image_url (base64) SELALU dibuang dari riwayat."""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "\n".join(
            p.get("text", "")
            for p in content
            if isinstance(p, dict) and p.get("type") == "text"
        )
    return ""

def trim_history(raw_history):
    """Bersihkan riwayat lalu ambil MAX_HISTORY_MESSAGES pesan terakhir(teks saja)"""
    if not isinstance(raw_history, list) or MAX_HISTORY_MESSAGES <= 0:
        return []

    clean = []
    for m in raw_history:
        if not isinstance(m, dict) or m.get("role") not in ("user", "assistant"):
            continue
        text = _content_to_text(m.get("content")).strip()
        if not text:
            continue
        if m["role"] == "assistant" and text.startswith(_NOISE_PREFIXES):
            #buang juga pertanyaan user yang menghasilkan error tersebut.
            if clean and clean[-1]["role"] == "user":
                clean.pop()
            continue
        if len(text) > MAX_HISTORY_CHARS_PER_MSG:
            text = text[:MAX_HISTORY_CHARS_PER_MSG] + "\n...(dipotong)..."
        clean.append({"role": m["role"], "content": text})

    window = clean[-MAX_HISTORY_MESSAGES:]
    while window and window[0]["role"] != "user":  #jendela harus diawali pesan user
        window.pop(0)
    return window

def enforce_total_file_budget(attachments):
    """Batasi total teks dokumen per request agar tidak menghabiskan TPM."""
    remaining = MAX_TOTAL_FILE_CHARS
    for a in attachments:
        if a["kind"] != "text":
            continue
        text = a["text"]
        if len(text) > remaining:
            text = text[:max(remaining, 0)] + "\n...(dipotong: total isi file melebihi batas per permintaan)..."
        a["text"] = text
        remaining = max(remaining - len(text), 0)
        
#CACHE IN-MEMORY(lapis pertama, sebelum memanggil Groq)
#catatan serverless:cache ini hidup per instance Vercel dan hilang saat cold start
_response_cache = TTLCache(maxsize=CACHE_MAX_ENTRIES, ttl=CACHE_TTL_SECONDS)
_cache_lock = threading.Lock()

def make_cache_key(model, question, history):
    """Kunci = model + pertanyaan (dinormalisasi) + sidik jari konteks riwayat.

    Untuk pesan pertama percakapan riwayatnya kosong, jadi kuncinya murni pertanyaan.
    Untuk pesan lanjutan ("jelaskan lebih detail", "ya"), konteks ikut masuk kunci
    agar tidak salah mengembalikan jawaban dari percakapan lain.
    """
    q = re.sub(r"\s+", " ", question.strip()).casefold()
    h = hashlib.sha256(
        "\x1f".join(f"{m['role']}:{m['content']}" for m in history).encode("utf-8")
    ).hexdigest() if history else ""
    return hashlib.sha256("\x1e".join([model, q, h]).encode("utf-8")).hexdigest()

def cache_get(key):
    with _cache_lock:
        return _response_cache.get(key)

def cache_set(key, answer):
    with _cache_lock:
        _response_cache[key] = answer

def sse(obj):
    return f"data: {json.dumps(obj)}\n\n"

@app.route("/api/chat", methods=["POST"])
def chat():
    if not client:
        return jsonify({"error": "GROQ_API_KEY belum diset di server."}), 500

    try:
        raw_payload = request.form.get("payload")
        if raw_payload:
            payload = json.loads(raw_payload)
        else:
            payload = request.get_json(force=True, silent=True) or {}
    except (ValueError, TypeError):
        return jsonify({"error": "Payload tidak valid."}), 400

    user_message = (payload.get("message") or "").strip()
    #hanya pesan terakhir, teks saja (tanpa base64 gambar)
    history = trim_history(payload.get("history"))

    attachments = []
    has_image_attachment = False
    for f in request.files.getlist("files"):
        filename = f.filename or "file"
        ext = os.path.splitext(filename)[1].lower()
        raw = f.read()

        if len(raw) > MAX_FILE_SIZE:
            return jsonify({"error": f"File '{filename}' terlalu besar (maks 15MB)."}), 400

        if ext in IMAGE_EXTS:
            mime = mimetypes.guess_type(filename)[0] or "image/png"
            attachments.append({
                "kind": "image",
                "name": filename,
                "mime": mime,
                "data_b64": base64.b64encode(raw).decode("utf-8"),
            })
            has_image_attachment = True
        elif ext == ".pdf":
            file_text = extract_text_from_pdf(raw)
            if len(file_text) > MAX_FILE_CHARS:
                file_text = file_text[:MAX_FILE_CHARS] + "\n...(dipotong, dokumen terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": file_text})
        elif ext in {".docx", ".doc"}:
            file_text = extract_text_from_docx(raw)
            if len(file_text) > MAX_FILE_CHARS:
                file_text = file_text[:MAX_FILE_CHARS] + "\n...(dipotong, dokumen terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": file_text})
        elif ext in {".xlsx", ".xls"}:
            file_text = extract_text_from_excel(raw)
            if len(file_text) > MAX_FILE_CHARS:
                file_text = file_text[:MAX_FILE_CHARS] + "\n...(dipotong, spreadsheet terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": file_text})
        elif ext in {".pptx", ".ppt"}:
            file_text = extract_text_from_pptx(raw)
            if len(file_text) > MAX_FILE_CHARS:
                file_text = file_text[:MAX_FILE_CHARS] + "\n...(dipotong, slide terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": file_text})
        elif ext == ".ipynb":
            file_text = extract_text_from_ipynb(raw)
            if len(file_text) > MAX_FILE_CHARS:
                file_text = file_text[:MAX_FILE_CHARS] + "\n...(dipotong, notebook terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": file_text})
        elif ext in CODE_AND_TEXT_EXTS:
            decoded = try_decode_as_text(raw) or "(gagal membaca isi file teks)"
            if len(decoded) > MAX_FILE_CHARS:
                decoded = decoded[:MAX_FILE_CHARS] + "\n...(dipotong, file terlalu panjang)..."
            attachments.append({"kind": "text", "name": filename, "text": decoded})
        else:
            decoded = try_decode_as_text(raw)
            if decoded is not None:
                if len(decoded) > MAX_FILE_CHARS:
                    decoded = decoded[:MAX_FILE_CHARS] + "\n...(dipotong, file terlalu panjang)..."
                attachments.append({"kind": "text", "name": filename, "text": decoded})
            else:
                attachments.append({
                    "kind": "text",
                    "name": filename,
                    "text": f"(tipe file '{ext}' berupa data biner yang belum didukung untuk dibaca langsung)",
                })

    if not user_message and not attachments:
        return jsonify({"error": "Pesan kosong."}), 400

    if not user_message and attachments:
        user_message = "Tolong analisis lampiran ini."

    enforce_total_file_budget(attachments)

    image_count = sum(1 for a in attachments if a["kind"] == "image")
    if image_count > MAX_IMAGES_PER_REQUEST:
        return jsonify({"error": f"Maksimal {MAX_IMAGES_PER_REQUEST} gambar per pesan."}), 400

    user_content, is_multimodal = build_user_content(user_message, attachments)

    #urutan payload dijaga agar prefix-nya stabil untuk Groq prompt caching
    #[system (statis, tanpa nilai dinamis)] -> [riwayat ter-trim, teks saja] -> [pesan user saat ini]
    #konten dinamis (isi file, base64 gambar) hanya ada di pesan user PALING AKHIR
    messages = [{"role": "system", "content": SYSTEM_PROMPT}]
    messages.extend(history)
    messages.append({"role": "user", "content": user_content})

    model_to_use = VISION_MODEL if has_image_attachment else TEXT_MODEL
    trigger_earth_image = (not has_image_attachment) and is_earth_image_intent(user_message)

    #cache hanya untuk pesan teks murni (tanpa lampiran) dan bukan permintaan gambar NASA.
    cache_key = None
    if not attachments and not trigger_earth_image:
        cache_key = make_cache_key(model_to_use, user_message, history)

    def generate():
        if trigger_earth_image:
            try:
                earth_payload = get_earth_today_payload()
                yield sse({'earth_image': earth_payload})
            except Exception as e:
                yield sse({'earth_image': {'ok': False, 'error': str(e)}})
            yield sse({'done': True})
            return

        #lapis 1:cek cache dulu, tidak ada HTTP request ke Groq jika hit
        if cache_key:
            cached = cache_get(cache_key)
            if cached is not None:
                for i in range(0, len(cached), 60):
                    yield sse({'token': cached[i:i + 60]})
                yield sse({'done': True, 'cached': True})
                return

        parts = []
        try:
            stream = client.chat.completions.create(
                model=model_to_use,
                messages=messages,
                temperature=0.7,
                max_tokens=MAX_OUTPUT_TOKENS,
                top_p=1,
                stream=True,
            )
            for chunk in stream:
                if not chunk.choices:
                    continue
                delta = chunk.choices[0].delta.content
                if delta:
                    parts.append(delta)
                    yield sse({'token': delta})

            #simpan ke cache hanya jika stream selesai penuh(bukan error/dihentikan user)
            full_answer = "".join(parts)
            if cache_key and full_answer.strip() and len(full_answer) <= CACHE_MAX_ANSWER_CHARS:
                cache_set(cache_key, full_answer)
            yield sse({'done': True})
        except RateLimitError as e:
            retry_after = None
            try:
                retry_after = e.response.headers.get("retry-after")
            except Exception:
                pass
            wait = f" Coba lagi sekitar {retry_after} detik." if retry_after else " Coba lagi beberapa saat lagi."
            yield sse({'error': "Batas permintaan Groq (rate limit) tercapai." + wait})
        except APIConnectionError:
            yield sse({'error': "Tidak bisa terhubung ke server Groq. Coba lagi."})
        except APIStatusError as e:
            if e.status_code == 413:
                yield sse({'error': "Permintaan terlalu besar untuk model. Kirim file yang lebih kecil atau mulai percakapan baru."})
            else:
                yield sse({'error': str(e)})
        except Exception as e:
            yield sse({'error': str(e)})

    return Response(
        stream_with_context(generate()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )
